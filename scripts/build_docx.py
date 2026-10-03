#!/usr/bin/env python3
"""최종 원고(md)를 Word 문서(.docx)로 변환하고, 변환 결과를 다시 열어 검사한다.

사용법:
  python3 build_docx.py <15_manuscript.md> --out <final.docx> [--images 폴더] [--toc 경로]
                        [--citations 01_citations.json] [--report build-report.md]
                        [--pdf] [--allow-missing-images] [--no-gate]

- 게이트: book.py의 publish_gate와 같은 판정이다 (모든 단계 완료, 설정 변경 없음, 병합 뒤 수정 기록,
  원고 재검사 🔴 0, 낮은 점수 승인, 사실 확인 승인). 출력 위치와 상관없이 원고가 있는 책 폴더의 상태로 판단하고,
  설정(--toc)과 출처(--citations)도 그 책의 것만 쓴다. 검사 없이 시험 삼아 변환할 때만 --no-gate를 쓰며,
  이때 결과물 머리글과 표지에 '시험 변환본'이 찍힌다
- --pdf: LibreOffice(soffice)가 있으면 PDF도 만들고 실제 쪽수를 목표 분량과 비교한다
- --allow-missing-images: 그림이 없어도 변환한다. 이때 판정은 '통과'가 아니라 '시험 변환'이다

서식은 user-book-toc.md의 "## 출판 설정"(키: 값)에서 읽고, 없는 키는 PUB_DEFAULTS를 쓴다.
원고 마크업 규약은 SKILL.md "원고 마크업 규약"을 따른다.

종료 코드: 0 통과, 1 변환 후 검사 실패, 2 사용법 오류
"""
import argparse
import datetime
import json
import os
import pathlib
import re
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import _common as C  # noqa: E402

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:  # doctor가 PUB_DEFAULTS만 읽을 수 있게 한다
    Document = None

PUB_DEFAULTS = {
    "판형": "B5",                  # B5, A5, A4, 신국판, 또는 '176x250' (mm)
    "본문 폰트": "Pretendard",
    "대체 폰트": "Malgun Gothic",   # 빌드 PC에 본문 폰트가 없을 때
    "본문 크기": "11",
    "줄간격": "1.2",
    "문단 뒤 간격": "6",
    "여백": "2.5",                 # cm, 상하좌우
    "주 색상": "#1F4E79",
    "보조 색상": "#2E75B6",
    "캡션 색상": "#666666",
    "박스 배경색": "#E8F0FE",
    "박스 테두리색": "#1F4E79",
    "그림 폭": "",                 # cm. 비우면 본문 폭의 90%
    "장 새 페이지": "예",
    "표지": "예",
    "목차 페이지": "예",
    "머리글": "예",                # 머리글에 책 제목
    "AI 제작 고지": "AI 도구의 도움을 받아 쓰고 사람이 검수한 책",  # 명사형이라 어미·화자와 무관. '없음'이면 넣지 않음
    "배포 표기": "없음",           # 예: '사내 한정', '대외비'. 표지와 모든 쪽 머리글에 찍힌다
}
# PUB_DEFAULTS는 user_input/user-book-toc.md 템플릿의 '출판 설정' 값과 같게 유지한다.
PAPER = C.PAPER
FIG_LINE_RE = re.compile(r"^\[그림\s*(\d+)\s*[:.]\s*(.*?)\]\s*$")
CITE_RE = re.compile(r"\[\[\s*(cite_[\w-]+(?:\s*,\s*cite_[\w-]+)*)\s*\]\]")
LIST_RE = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")


def rgb(hexstr):
    h = hexstr.strip().lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def font_available(name):
    dirs = [pathlib.Path("/Library/Fonts"), pathlib.Path.home() / "Library/Fonts", pathlib.Path("/System/Library/Fonts"),
            pathlib.Path("C:/Windows/Fonts"), pathlib.Path("/usr/share/fonts"), pathlib.Path.home() / ".fonts"]
    key = re.sub(r"\s", "", name).lower()
    for d in dirs:
        if d.is_dir():
            for f in d.rglob("*"):
                if key in re.sub(r"[\s_-]", "", f.name).lower():
                    return True
    return False


class Builder:
    def __init__(self, toc_text, images_dir, bullet, trial=False):
        self.trial = trial
        info = C.basic_info(toc_text)
        vals, _ = C.parse_kv(C.get_section(toc_text, "출판 설정"), PUB_DEFAULTS.keys())
        self.cfg = dict(PUB_DEFAULTS)
        self.cfg.update({k: v for k, v in vals.items() if v})
        if self.cfg["AI 제작 고지"].strip() in ("없음", "-"):
            self.cfg["AI 제작 고지"] = ""
        if self.cfg["배포 표기"].strip() in ("없음", "-"):
            self.cfg["배포 표기"] = ""
        self.title = info.get("제목", "")
        self.subtitle = info.get("부제", "")
        self.author = info.get("필명", "")
        self.images_dir = pathlib.Path(images_dir) if images_dir else None
        self.bullet = bullet
        rules, _ = C.parse_kv(C.get_section(toc_text, "검증 규칙"), ["부 브릿지 문구"])
        self.bridge = rules.get("부 브릿지 문구", "")
        self.notes = []
        self.missing_images = []
        self.inserted_images = 0
        self.cite_num = {}

        self.font = self.cfg["본문 폰트"]
        if not font_available(self.font) and self.cfg["대체 폰트"]:
            self.notes.append(f"빌드 PC에 '{self.font}'가 없어 '{self.cfg['대체 폰트']}'로 대체함")
            self.font = self.cfg["대체 폰트"]
        self.size = float(self.cfg["본문 크기"])
        self.spacing = float(self.cfg["줄간격"])
        self.after = float(self.cfg["문단 뒤 간격"])
        self.primary = self.cfg["주 색상"]
        self.secondary = self.cfg["보조 색상"]
        self.doc = Document()
        self._setup()

    # ------------------------------------------------------------- 문서 설정
    def _setup(self):
        doc = self.doc
        paper = self.cfg["판형"].strip()
        if paper in PAPER:
            w, h = PAPER[paper]
        else:
            m = re.match(r"(\d+)\s*[x×]\s*(\d+)", paper)
            w, h = (int(m.group(1)), int(m.group(2))) if m else PAPER["B5"]
        margin = float(self.cfg["여백"])
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(w / 10), Cm(h / 10)
        for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(sec, side, Cm(margin))
        self.text_width_cm = w / 10 - 2 * margin
        self.page_height_cm = h / 10
        self.fig_width = float(self.cfg["그림 폭"]) if self.cfg["그림 폭"] else round(self.text_width_cm * 0.9, 1)

        # 기본 스타일: 한글 폰트, 언어, 한글과 영문·숫자 사이 자동 간격 끄기
        normal = doc.styles["Normal"]
        self._style_font(normal, self.size)
        for rf in doc.styles.element.iter(qn("w:rFonts")):  # 문서 기본값의 테마 글꼴도 지운다
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(attr)) is not None:
                    del rf.attrib[qn(attr)]
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
                if rf.get(qn(attr)) is None:
                    rf.set(qn(attr), self.font)
        pf = normal.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(self.after), self.spacing
        ppr = normal.element.get_or_add_pPr()
        ppr.append(parse_xml(f'<w:autoSpaceDE {nsdecls("w")} w:val="0"/>'))
        ppr.append(parse_xml(f'<w:autoSpaceDN {nsdecls("w")} w:val="0"/>'))
        ppr.append(parse_xml(f'<w:wordWrap {nsdecls("w")} w:val="0"/>'))
        rpr = normal.element.get_or_add_rPr()
        rpr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="ko-KR" w:eastAsia="ko-KR"/>'))

        for name, size, color, before, after in [
            ("Heading 1", 22, self.primary, 0, 8),
            ("Heading 2", 18, self.primary, 0, 10),
            ("Heading 3", 14, self.primary, 14, 6),
            ("Heading 4", 12, self.secondary, 10, 4),
        ]:
            st = doc.styles[name]
            self._style_font(st, size, bold=True, color=color)
            st.font.italic = False  # 기본 템플릿의 Heading 3·4는 기울임이다
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.keep_with_next = True

        sec.different_first_page_header_footer = True  # 표지에는 머리글·쪽 번호를 넣지 않는다
        if (C.is_yes(self.cfg["머리글"]) and self.title) or self.trial or self.cfg["배포 표기"]:
            hp = sec.header.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            label = ("시험 변환본 (검사·승인 전) · " if self.trial else "") + (f"{self.cfg['배포 표기']} · " if self.cfg["배포 표기"] else "") \
                + (self.title if C.is_yes(self.cfg["머리글"]) else "")
            self._run(hp, label.strip(" ·"), 9, color="#C00000" if self.trial else self.cfg["캡션 색상"])
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._field(fp, "PAGE", 9)

        settings = doc.settings.element
        settings.append(parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>'))
        doc.core_properties.author = self.author
        doc.core_properties.comments = ""
        doc.core_properties.created = doc.core_properties.modified = build_time()
        doc.core_properties.subject = self.subtitle
        doc.core_properties.title = f"{self.title}: {self.subtitle}" if self.subtitle else self.title

    def _style_font(self, style, size, bold=False, color=None):
        style.font.name = self.font
        style.font.size = Pt(size)
        style.font.bold = bold
        if color:
            style.font.color.rgb = rgb(color)
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), self.font)
        # 테마 글꼴 속성이 있으면 명시한 글꼴보다 우선한다 (기본 템플릿의 제목이 Calibri Light 등으로 바뀌는 원인)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]

    def _run(self, p, text, size=None, bold=None, italic=None, color=None, font=None, sup=False):
        r = p.add_run(text)
        if size:
            r.font.size = Pt(size)
        if bold is not None:
            r.font.bold = bold
        if italic:
            r.font.italic = True
        if color:
            r.font.color.rgb = rgb(color)
        if sup:
            r.font.superscript = True
        if font:
            r.font.name = font
            rpr = r._element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
                rpr.insert(0, rf)
            rf.set(qn("w:ascii"), font)
            rf.set(qn("w:hAnsi"), font)
            rf.set(qn("w:eastAsia"), self.font)
        return r

    def _field(self, p, instr, size, placeholder="1"):
        r = p.add_run()
        r._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
        r = p.add_run()
        r._element.append(parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> {instr} </w:instrText>'))
        r = p.add_run()
        r._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>'))
        self._run(p, placeholder, size, color=None if placeholder == "1" else self.cfg["캡션 색상"])
        r = p.add_run()
        r._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))

    # ------------------------------------------------------------- 인라인
    def inline(self, p, text, size=None, color=None, bold=False):
        """**굵게**, *기울임*, `코드`, [링크](url), [[cite]], <br> 처리."""
        text = text.replace("[[예시]]", "").replace(" .", ".")
        text = re.sub(r"\s+(\[\[\s*cite_)", r"\1", text)  # 위첨자 번호 앞 공백 제거 ('3개입니다 ¹.' 방지)
        text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1", text)
        text = re.sub(r"\[그림\s*(\d+)\]", r"그림 \1", text)
        tokens = C.INLINE_SPLIT_RE.split(text)
        for tok in tokens:
            if not tok:
                continue
            if tok.startswith("***") and tok.endswith("***") and len(tok) > 6:
                self._run(p, tok[3:-3], size, bold=True, italic=True, color=color)
            elif (tok.startswith("**") and tok.endswith("**") and len(tok) > 4) or (tok.startswith("__") and tok.endswith("__") and len(tok) > 4):
                for part in re.split(r"(\*[^*\s][^*]*\*)", tok[2:-2]):  # 굵은 글 안의 *기울임*
                    if part.startswith("*") and part.endswith("*") and len(part) > 2:
                        self._run(p, part[1:-1], size, bold=True, italic=True, color=color)
                    elif part:
                        self._run(p, part, size, bold=True, color=color)
            elif tok.startswith("`") and tok.endswith("`"):
                self._run(p, tok[1:-1], size, color=color, font="Consolas")
            elif CITE_RE.fullmatch(tok):
                ids = [x.strip() for x in CITE_RE.fullmatch(tok).group(1).split(",")]
                nums = []
                for cid in ids:
                    if cid not in self.cite_num:
                        self.cite_num[cid] = len(self.cite_num) + 1
                    nums.append(str(self.cite_num[cid]))
                self._run(p, f"[{','.join(nums)}]", size, color=color, sup=True)
            elif re.fullmatch(r"<br\s*/?>", tok):
                p.add_run().add_break()
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                self._run(p, tok[1:-1], size, italic=True, color=color, bold=bold or None)
            else:
                tok = re.sub(r"</?[a-zA-Z][^>]*>", "", tok)
                self._run(p, tok, size, bold=bold or None, color=color)

    # ------------------------------------------------------------- 블록
    def para(self, text, **kw):
        p = self.doc.add_paragraph()
        self.inline(p, text, **kw)
        return p

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def heading(self, level, text):
        p = self.doc.add_paragraph(style=f"Heading {min(level, 4)}")
        if level == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_before = Pt(120)
        if level == 2 and C.is_yes(self.cfg["장 새 페이지"]):
            p.paragraph_format.page_break_before = True
        self.inline(p, text)
        return p

    def subtitle_line(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.inline(p, text, size=14, color=self.secondary)
        p.paragraph_format.space_after = Pt(36)

    def bridge_para(self, text):
        p = self.doc.add_paragraph()
        self._run(p, text, 13, bold=True, color=self.primary)
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.keep_with_next = True

    def list_item(self, indent, marker, text, size=None):
        p = self.doc.add_paragraph()
        level = indent // 2
        p.paragraph_format.left_indent = Cm(0.6 + 0.6 * level)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        p.paragraph_format.space_after = Pt(max(2, self.after / 2))
        sym = self.bullet if marker in "-*+" else marker
        self._run(p, sym + " ", size)
        self.inline(p, text, size=size)
        return p

    def quote(self, lines):
        p = self.doc.add_paragraph()
        ppr = p._element.get_or_add_pPr()
        ppr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="{self.secondary.lstrip("#")}"/></w:pBdr>'))
        ppr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F6F8FA"/>'))
        p.paragraph_format.left_indent = Cm(0.5)
        self.inline(p, "<br>".join(lines), size=self.size - 0.5, color="#333333")

    def code(self, lines):
        p = self.doc.add_paragraph()
        ppr = p._element.get_or_add_pPr()
        ppr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F3F4F6"/>'))
        ppr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="4" w:color="D0D7DE"/><w:left w:val="single" w:sz="4" w:space="4" w:color="D0D7DE"/><w:bottom w:val="single" w:sz="4" w:space="4" w:color="D0D7DE"/><w:right w:val="single" w:sz="4" w:space="4" w:color="D0D7DE"/></w:pBdr>'))
        p.paragraph_format.left_indent = Cm(0.2)
        p.paragraph_format.right_indent = Cm(0.2)
        p.paragraph_format.line_spacing = 1.1
        for i, line in enumerate(lines):
            if i:
                p.add_run().add_break()
            self._run(p, line, self.size - 1.5, font="Consolas")

    def figure(self, num, caption):
        img = None
        if self.images_dir:
            for name in (f"fig{num:02d}.png", f"fig{num}.png", f"fig{num:02d}.jpg"):
                if (self.images_dir / name).is_file():
                    img = self.images_dir / name
                    break
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(8)
        if img:
            p.add_run().add_picture(str(img), width=Cm(self.fig_width))
            self.inserted_images += 1
        else:
            self.missing_images.append(num)
            ppr = p._element.get_or_add_pPr()
            ppr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="dashed" w:sz="6" w:space="12" w:color="999999"/><w:left w:val="dashed" w:sz="6" w:space="12" w:color="999999"/><w:bottom w:val="dashed" w:sz="6" w:space="12" w:color="999999"/><w:right w:val="dashed" w:sz="6" w:space="12" w:color="999999"/></w:pBdr>'))
            self._run(p, f"이미지 없음: fig{num:02d}.png", 10, color="#999999")
        cp = self.doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(10)
        self.inline(cp, f"그림 {num}. {caption}", size=self.size - 1, color=self.cfg["캡션 색상"])

    def table(self, rows):
        header, body = rows[0], rows[1:]
        ncol = len(header)
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        # 열 너비를 내용 길이에 비례해 나눈다 (좁은 열의 글자가 한 자씩 끊기지 않게, 열마다 최소 1.6cm)
        lens = []
        for ci in range(ncol):
            cells = [re.sub(r"<[^>]+>|\*\*", "", r[ci]) if ci < len(r) else "" for r in rows]
            lens.append(max(4, min(40, max(len(x) for x in cells))))
        widths = [max(1.6, self.text_width_cm * x / sum(lens)) for x in lens]
        scale = self.text_width_cm / sum(widths)
        widths = [w * scale for w in widths]
        for ci, w in enumerate(widths):
            for cell in t.columns[ci].cells:
                cell.width = Cm(w)
        # 정말 짧은 표만 한 페이지에 붙인다. 셀 글이 긴 표까지 붙이면 렌더러가 빈 페이지를 만든다
        keep = len(rows) <= 9 and sum(len(c) for r in rows for c in r) < 600
        for ri, row in enumerate(rows):
            tr = t.rows[ri]._tr
            trpr = tr.get_or_add_trPr()
            if ri == 0:
                trpr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            trpr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            for ci in range(ncol):
                cell = t.cell(ri, ci)
                txt = row[ci] if ci < len(row) else ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if keep and ri < len(rows) - 1:
                    p.paragraph_format.keep_with_next = True
                if ri == 0:
                    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{self.primary.lstrip("#")}"/>'))
                    self.inline(p, txt, size=self.size - 1.5, color="#FFFFFF", bold=True)
                else:
                    segs = re.split(r"<br\s*/?>", txt)
                    for si, seg in enumerate(segs):
                        small = seg.strip().startswith("<sm>")
                        seg = re.sub(r"</?sm>", "", seg)
                        if si:
                            p.add_run().add_break()
                        self.inline(p, seg, size=8 if small else self.size - 1.5, color="#555555" if small else None)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)

    def box(self, lines):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.cell(0, 0)
        tcpr = cell._tc.get_or_add_tcPr()
        bc = self.cfg["박스 테두리색"].lstrip("#")
        tcpr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="8" w:color="{bc}"/><w:left w:val="single" w:sz="8" w:color="{bc}"/><w:bottom w:val="single" w:sz="8" w:color="{bc}"/><w:right w:val="single" w:sz="8" w:color="{bc}"/></w:tcBorders>'))
        tcpr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{self.cfg["박스 배경색"].lstrip("#")}"/>'))
        tcpr.append(parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="240" w:type="dxa"/><w:left w:w="240" w:type="dxa"/><w:bottom w:w="240" w:type="dxa"/><w:right w:w="240" w:type="dxa"/></w:tcMar>'))
        t.rows[0]._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        first = True
        content = [ln for ln in lines if ln.strip()]
        for i, ln in enumerate(content):
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.paragraph_format.space_after = Pt(3)
            if i < len(content) - 1:
                p.paragraph_format.keep_with_next = True
            s = ln.strip()
            m = LIST_RE.match(ln)
            if m:
                sym = self.bullet if m.group(2) in "-*+" else m.group(2)
                p.paragraph_format.left_indent = Cm(0.45)
                p.paragraph_format.first_line_indent = Cm(-0.45)
                self._run(p, sym + " ")
                self.inline(p, m.group(3))
            elif i == 0 or s.startswith("#"):
                self.inline(p, s.lstrip("#").strip().strip("*"), size=self.size + 1, color=self.primary, bold=True)
            else:
                self.inline(p, s)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)

    def cover(self):
        # 빈 줄 대신 판형 높이에 비례한 간격을 둔다 (작은 판형에서 표지가 넘쳐 빈 쪽이 생기지 않게)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Cm(self.page_height_cm * 0.22)
        self._run(p, self.title, 28, bold=True, color=self.primary)
        if self.subtitle:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._run(p, self.subtitle, 16, color=self.secondary)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Cm(self.page_height_cm * 0.08)
        self._run(p, self.author, 14, color=self.primary)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        today = build_time()
        self._run(p, f"{today.year}년 {today.month}월", 11, color=self.cfg["캡션 색상"])
        if self.cfg["배포 표기"]:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._run(p, self.cfg["배포 표기"], 12, bold=True, color="#C00000")
        if self.cfg["AI 제작 고지"]:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._run(p, self.cfg["AI 제작 고지"], 9, color=self.cfg["캡션 색상"])
        if self.trial:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._run(p, "시험 변환본: 규칙 검사와 사실 확인을 거치지 않았습니다. 배포하지 마세요.", 10, bold=True, color="#C00000")

    def toc_page(self, md, page_of=None):
        """차례. Word의 TOC 필드를 넣고, 필드가 갱신되지 않는 환경(PDF 변환 등)을 위해
        필드 결과 자리에 제목 목록을 미리 채운다. Word에서 필드를 갱신하면 쪽 번호가 붙은 차례로 바뀐다."""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.page_break_before = len(self.doc.paragraphs) > 1  # 표지가 있으면 새 쪽에서
        self._run(p, "차례", 20, bold=True, color=self.primary)
        entries = [(lv, re.sub(r"<br\s*/?>", " ", t)) for _, lv, t in C.iter_headings(md) if lv <= 2]
        bridge_re = None
        if self.bridge and "N" in self.bridge:
            bridge_re = re.compile("^" + re.escape(self.bridge).replace("N", r"(제\s*)?\d+") + r"\s*$")
        entries = [(lv, t) for lv, t in entries if not (lv == 2 and bridge_re and bridge_re.match(t))]
        if re.search(r"\[\[\s*cite_", md):
            entries.append((1, "출처"))
        first = self.doc.add_paragraph()
        r = first.add_run()
        r._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
        r = first.add_run()
        r._element.append(parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> TOC \\o "1-2" \\h \\z \\u </w:instrText>'))
        r = first.add_run()
        r._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>'))
        para = first
        self.toc_entries = entries
        for i, (lv, t) in enumerate(entries or [(1, "")]):
            if i:
                para = self.doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0 if lv == 1 else 0.8)
            para.paragraph_format.space_after = Pt(2 if lv == 2 else 6)
            para.paragraph_format.tab_stops.add_tab_stop(Cm(self.text_width_cm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            self._run(para, re.sub(r"\*\*|`", "", t), 11 if lv == 1 else 10, bold=lv == 1, color=self.primary if lv == 1 else None)
            if page_of and page_of.get(i):
                self._run(para, f"\t{page_of[i]}", 10)
        r = para.add_run()
        r._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))

    # ------------------------------------------------------------- 본문 변환
    def render(self, md):
        lines = md.splitlines()
        i, n = 0, len(lines)
        last_h1 = False
        bridge_re = None
        if self.bridge and "N" in self.bridge:
            bridge_re = re.compile("^" + re.escape(self.bridge).replace("N", r"(제\s*)?\d+") + r"\s*$")
        while i < n:
            line = lines[i]
            s = line.strip()
            if not s:
                i += 1
                continue
            if s.startswith(("```", "~~~")):
                fence = s[:3]
                j = i + 1
                block = []
                while j < n and not lines[j].strip().startswith(fence):
                    block.append(lines[j])
                    j += 1
                self.code(block)
                i = j + 1
                continue
            if s == "<!-- box-start -->":
                j = i + 1
                block = []
                while j < n and lines[j].strip() != "<!-- box-end -->":
                    block.append(lines[j])
                    j += 1
                self.box(block)
                i = j + 1
                continue
            if s == "<!-- pagebreak -->":
                # 다음 블록이 어차피 새 페이지에서 시작하는 제목이면 빈 페이지가 생기므로 건너뛴다
                j = i + 1
                while j < n and not lines[j].strip():
                    j += 1
                nxt = lines[j].strip() if j < n else ""
                if not (nxt.startswith("# ") or (nxt.startswith("## ") and C.is_yes(self.cfg["장 새 페이지"]))):
                    self.page_break()
                i += 1
                continue
            if s.startswith("<!--"):
                while i < n and "-->" not in lines[i]:
                    i += 1
                i += 1
                continue
            if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", s):
                i += 1
                continue
            m = re.match(r"^(#{1,6})\s+(.*)$", s)
            if m:
                level, text = len(m.group(1)), m.group(2).strip()
                if level == 2 and bridge_re and bridge_re.match(re.sub(r"\s+", " ", text)):
                    self.bridge_para(text)
                else:
                    self.heading(level, text)
                last_h1 = level == 1
                i += 1
                if last_h1:
                    j = i
                    while j < n and not lines[j].strip():
                        j += 1
                    if j < n:
                        nxt = lines[j].strip()
                        sub = None
                        if re.fullmatch(r"\*\*[^*]+\*\*", nxt):
                            sub = nxt[2:-2]
                        elif nxt.startswith("부제:"):
                            sub = nxt[3:].strip().strip('"“”')
                        if sub is not None:
                            self.subtitle_line(sub)
                            i = j + 1
                continue
            fm = FIG_LINE_RE.match(s)
            if fm:
                self.figure(int(fm.group(1)), fm.group(2).strip())
                i += 1
                continue
            if s.startswith("|"):
                rows = []
                while i < n and lines[i].strip().startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                        rows.append(cells)
                    i += 1
                if rows:
                    self.table(rows)
                continue
            if s.startswith(">"):
                block = []
                while i < n and lines[i].strip().startswith(">"):
                    block.append(lines[i].strip()[1:].strip())
                    i += 1
                self.quote(block)
                continue
            lm = LIST_RE.match(line)
            if lm:
                self.list_item(len(lm.group(1)), lm.group(2), lm.group(3))
                i += 1
                continue
            # 일반 문단: 빈 줄 전까지 이어 붙인다
            block = [s]
            i += 1
            while i < n and lines[i].strip() and not re.match(r"^(#|\||>|```|<!--|\[그림)", lines[i].strip()) and not LIST_RE.match(lines[i]):
                block.append(lines[i].strip())
                i += 1
            self.para(" ".join(block))

    def references(self, cites):
        if not self.cite_num:
            return
        self.heading(1, "출처")
        for cid, num in sorted(self.cite_num.items(), key=lambda x: x[1]):
            line = f"[{num}] {C.citation_line(cites.get(cid, {'id': cid}))}"
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            self._run(p, line, self.size - 2)


def post_check(docx_path, md, allow_emdash, builder, allow_missing, cites):
    """변환된 docx를 다시 열어 남으면 안 되는 것을 찾는다."""
    doc = Document(str(docx_path))
    texts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs)
    full = "\n".join(texts)
    # 코드 블록과 인라인 코드 안의 글자는 그대로 실리는 것이 정상이다. 그만큼은 허용한다.
    code = "\n".join([m[1] for m in re.findall(r"^(```|~~~)[^\n]*\n(.*?)^\1", md, re.M | re.S)] + re.findall(r"`([^`\n]+)`", md))
    fails, warns = [], []
    for pat, label in [(r"\*\*", "마크다운 ** 잔존"), (r"(?<![\w_])__|__(?![\w_])", "마크다운 __ 잔존"), (r"<!--", "HTML 주석 잔존"), (r"\[\[cite", "출처 태그 미변환"),
                       (r"<br|<sm>", "HTML 태그 잔존"), (r"\[그림\s*\d+\s*[:.]", "그림 자리표시 미변환")]:
        k = len(re.findall(pat, full)) - len(re.findall(pat, code))
        if k > 0:
            fails.append(f"{label} {k}건")
    lone = r"(?<![*\w])\*(?![*\s])|(?<![*\s])\*(?![*\w])"
    k = len(re.findall(lone, full)) - len(re.findall(lone, code))
    if k > 0:
        warns.append(f"홀로 남은 * {k}건 (강조 표시가 짝을 이루지 않았을 수 있다)")
    k = full.count("~~") - code.count("~~")
    if k > 0:
        warns.append(f"취소선 기호 ~~ {k}건 (Word에 그대로 찍힌다)")
    if not allow_emdash:
        k = len(re.findall(r"[—–]", full))
        if k:
            fails.append(f"em dash {k}건")
    h1 = sum(1 for p in doc.paragraphs if p.style.name == "Heading 1")
    h2 = sum(1 for p in doc.paragraphs if p.style.name == "Heading 2")
    md_h1 = sum(1 for _, lv, _ in C.iter_headings(md) if lv == 1) + (1 if builder.cite_num else 0)  # 책 끝 '출처
    if h1 != md_h1:
        fails.append(f"부 헤딩 수 불일치 (원고 {md_h1}, docx {h1})")
    if builder.missing_images:
        miss = builder.missing_images
        msg = f"그림 파일 없음 {len(miss)}개: " + ", ".join(f"fig{x:02d}" for x in miss[:20]) + (" …" if len(miss) > 20 else "")
        (warns if allow_missing else fails).append(msg)
    unknown = [cid for cid in builder.cite_num if cid not in cites]
    if unknown:
        fails.append(f"출처 파일에 없는 출처 ID {len(unknown)}개: {', '.join(unknown[:10])}")
    if doc.core_properties.author != builder.author:
        fails.append("문서 작성자 속성이 필명과 다름")
    return fails, warns, {"h1": h1, "h2": h2, "images": builder.inserted_images}


def main(argv=None):
    ap = argparse.ArgumentParser(description="원고 → docx 변환과 변환 후 검사")
    ap.add_argument("manuscript")
    ap.add_argument("--out", required=True)
    ap.add_argument("--images")
    ap.add_argument("--toc")
    ap.add_argument("--citations")
    ap.add_argument("--report")
    ap.add_argument("--allow-missing-images", action="store_true")
    ap.add_argument("--pdf", action="store_true")
    ap.add_argument("--no-gate", action="store_true")
    a = ap.parse_args(argv)
    if Document is None:
        print(f"python-docx가 없습니다: pip install -r \"{C.SKILL_DIR / 'scripts' / 'requirements.txt'}\"", file=sys.stderr)
        return 2
    src = pathlib.Path(a.manuscript)
    toc_path = pathlib.Path(a.toc) if a.toc else C.find_toc()
    if not src.is_file() or not toc_path or not toc_path.is_file():
        print("원고 또는 책 설정 파일을 찾지 못했습니다.", file=sys.stderr)
        return 2
    toc = toc_path.read_text(encoding="utf-8")
    md = C.strip_markers(src.read_text(encoding="utf-8"))
    rules, _ = C.parse_kv(C.get_section(toc, "검증 규칙"), ["불릿 기호", "em dash 허용"])
    bullet = (rules.get("불릿 기호") or "•").strip()[0]
    out = pathlib.Path(a.out)
    images = a.images or str(out.parent / "images")
    cite_path = pathlib.Path(a.citations) if a.citations else src.resolve().with_name("01_citations.json")  # 링크를 따라간 실제 책 폴더
    cites = {}
    if cite_path.is_file():
        try:
            cites = {c["id"]: c for c in json.loads(cite_path.read_text(encoding="utf-8"))}
        except (ValueError, KeyError, TypeError) as e:
            print(f"출처 파일을 읽지 못했습니다 ({cite_path.name}: {e}). JSON 목록이고 항목마다 id가 있어야 합니다.", file=sys.stderr)
            return 2

    if not a.no_gate:
        reason = gate_reason(src, out)
        if reason:
            print(f"변환 중단: {reason} (검사 없이 시험 삼아 변환할 때만 --no-gate)", file=sys.stderr)
            return 1
        # 게이트를 통과한 변환은 그 책의 설정과 출처만 쓴다 (다른 설정 파일로 고지나 규칙을 바꿔 끼우지 못하게)
        work = src.resolve().parent.parent.parent
        book_toc = C.input_dir(work) / C.TOC_NAME
        if a.toc and pathlib.Path(a.toc).resolve() != book_toc.resolve():
            print("변환 중단: 출판 변환에는 이 책의 설정(user_input/user-book-toc.md)만 쓸 수 있습니다. --toc를 빼세요.", file=sys.stderr)
            return 1
        if a.citations and pathlib.Path(a.citations).resolve() != (src.resolve().parent / "01_citations.json").resolve():
            print("변환 중단: 출판 변환에는 이 책의 01_citations.json만 쓸 수 있습니다. --citations를 빼세요.", file=sys.stderr)
            return 1
        toc = book_toc.read_text(encoding="utf-8")
        book_images = (work / "output" / src.resolve().parent.name / "images").resolve()
        if a.images and pathlib.Path(a.images).resolve() != book_images:
            print("변환 중단: 출판 변환에는 이 책의 그림 폴더(output/{책}/images)만 쓸 수 있습니다. --images를 빼세요.", file=sys.stderr)
            return 1
        images = str(book_images)

    def build(page_of=None):
        bld = Builder(toc, images, bullet, trial=a.no_gate)
        if C.is_yes(bld.cfg["표지"]):
            bld.cover()
        if C.is_yes(bld.cfg["목차 페이지"]):
            bld.toc_page(md, page_of)
        bld.render(md)
        bld.references(cites)
        out.parent.mkdir(parents=True, exist_ok=True)
        bld.doc.save(str(out))
        return bld

    if out.with_suffix(".pdf").exists():
        out.with_suffix(".pdf").unlink()  # 이전 빌드의 PDF가 새 결과처럼 남지 않게 (--pdf가 없어도 지운다)
    b = build()
    fails, warns, stats = post_check(out, md, C.is_yes(rules.get("em dash 허용", "")), b, a.allow_missing_images, cites)
    # 마지막 관문: 완성된 문서의 글자 전체(본문, 코드 블록, 표, 머리글, 출처 목록)와 그림 기록을 verify와 같은 기준으로 다시 본다
    toc_file = (C.input_dir(src.resolve().parent.parent.parent) / C.TOC_NAME) if not a.no_gate else toc_path
    lf, lw = leak_check(out, toc, toc_file, src, b.author)
    fails += [] if a.no_gate else lf
    warns += lw + (lf if a.no_gate else [])  # 시험 변환은 쪽수를 재는 용도라 경고로만 남긴다 (출판 변환에서는 실패)
    pages = render_pdf(out) if a.pdf else None
    if pages and getattr(b, "toc_entries", None):
        # 2단계: PDF에서 각 제목의 실제 쪽을 찾아 차례에 넣고 다시 렌더링한다 (Word 필드 갱신 없이도 쪽 번호가 보이게)
        for _ in range(3):
            page_of = find_heading_pages(out.with_suffix(".pdf"), b.toc_entries)
            if not page_of:
                break
            b = build(page_of)
            pages = render_pdf(out)
            if find_heading_pages(out.with_suffix(".pdf"), b.toc_entries) == page_of:
                break
        else:
            # 쪽 번호가 렌더링마다 달라지면 틀린 번호를 찍지 않고 차례에서 뺀다 (Word에서 '필드 업데이트'로 채울 수 있다)
            b = build(None)
            pages = render_pdf(out)
            warns.append("차례 쪽 번호가 렌더링마다 달라져 차례에서 쪽 번호를 뺐다 (Word에서 '필드 업데이트'로 채운다)")
    if pages:
        info = C.basic_info(toc).get("목표 분량", "")
        m = re.search(r"(\d+)\s*(?:~\s*(\d+)\s*)?(?:페이지|쪽)", info)
        if m:
            lo = int(m.group(1))
            hi = int(m.group(2) or m.group(1))
            if pages < lo * 0.9 or pages > hi * 1.15:
                warns.append(f"실측 쪽수 {pages}쪽이 목표 {m.group(0)}에서 벗어남 (LibreOffice 렌더링 기준)")
    verdict = "실패" if fails else ("시험 변환 (게이트 없음)" if a.no_gate else "시험 변환 (그림 누락 허용)" if b.missing_images else "통과")
    report = pathlib.Path(a.report) if a.report else out.with_name("build-report.md")
    lines = ["# build-report", "", f"- 원고: `{src.name}`", f"- 원고 해시: {C.sha(src)}", f"- 결과: `{out.name}`", f"- 결과 해시: {C.sha(out)}",
             f"- 부 {stats['h1'] - (1 if b.cite_num else 0)}개, 장 {stats['h2']}개, 그림 {stats['images']}개 삽입",
             f"- 실측 쪽수: {pages}쪽 (PDF: {out.with_suffix('.pdf').name})" if pages else "- 실측 쪽수: 측정 안 함 (--pdf와 LibreOffice 필요)",
             f"- 판정: {verdict}", ""]
    lines += [f"- 🔴 {x}" for x in fails] + [f"- 🟡 {x}" for x in warns] + [f"- ℹ️ {x}" for x in b.notes]
    lines += ["", "차례: Word로 열 때 '필드 업데이트'를 허용하면 쪽 번호가 붙는다. 갱신 전에도 제목 목록은 들어 있다."]
    report.write_text("\n".join(lines) + "\n", encoding="utf-8")
    if a.pdf and not pages:
        print("사용자 안내: 이 환경에는 LibreOffice가 없어 PDF와 실제 쪽수를 만들지 못했습니다. final.docx를 Word(또는 한컴오피스, Google 문서)로 열어 "
              "'필드 업데이트'를 허용한 뒤 'PDF로 저장'하면 배포용 PDF가 됩니다. 마지막 쪽 번호가 실제 쪽수입니다.")
    print(f"docx: {out} | 부 {stats['h1'] - (1 if b.cite_num else 0)}, 장 {stats['h2']}, 그림 {stats['images']}" + (f", {pages}쪽" if pages else "")
          + f" | {verdict}" + ("" if not fails else ": " + "; ".join(fails)))
    return 0 if not fails else 1


def leak_check(docx_path, toc, toc_file, src, author):
    """완성된 docx에 새면 안 되는 것(금지 용어, 쓰면 안 되는 것, 실명)이 남았는지. (실패 목록, 경고 목록)"""
    import verify
    doc = Document(str(docx_path))
    texts = [p.text for p in doc.paragraphs]
    texts += [p.text for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
    texts += [p.text for s in doc.sections for p in s.header.paragraphs]
    rules = dict(verify.RULE_DEFAULTS)
    rules.update({k: v for k, v in C.parse_kv(C.get_section(toc, "검증 규칙"), verify.RULE_DEFAULTS.keys())[0].items() if v})
    toc_file = pathlib.Path(toc_file)
    sdir = toc_file.resolve().parent / "sources"
    files = [x for x in sdir.rglob("*") if x.is_file() and not x.name.startswith(".")] if sdir.is_dir() else []
    files += [x for x in [src.resolve().parent / "01_research-notes.md"] if x.is_file()]
    P = verify.leak_policy(toc, rules, files)
    P["allow"].add(C.despace(author))  # 표지의 필명은 저자가 정한 이름이다
    P["src_names"].discard(C.despace(author))
    blobs = ["\n".join(texts)] + [t for _, t in verify.figure_texts(toc_file, src, src.read_text(encoding="utf-8"))]
    hits = [h for blob in blobs for h in verify.leak_scan(blob, P)]
    fails = sorted({f"{cat}: {re.sub(r' [(].*$', '', detail)}" for sev, cat, detail, _ in hits if sev == "🔴"})
    warns = sorted({f"{cat}: {re.sub(r' [(].*$', '', detail)}" for sev, cat, detail, _ in hits if sev == "🟡"})
    return ([f"완성 문서에 새면 안 되는 표현 {len(fails)}건: " + "; ".join(fails[:8])] if fails else []), \
        ([f"완성 문서에서 확인할 표현 {len(warns)}건: " + "; ".join(warns[:8])] if warns else [])


def build_time():
    """문서에 찍는 시각. SOURCE_DATE_EPOCH가 있으면 그 시각을 써서 같은 원고는 같은 문서가 되게 한다."""
    epoch = os.environ.get("SOURCE_DATE_EPOCH")
    if epoch and epoch.isdigit():
        return datetime.datetime.utcfromtimestamp(int(epoch))
    return datetime.datetime.now()


def gate_reason(src, out, toc_text=None):
    """변환해도 되는지. 안 되면 이유 문자열.

    판정은 book.py의 publish_gate 하나로 한다 (status, approve와 같은 기준).
    원고가 책 폴더(draft/{책}/15_manuscript.md)에 있고 상태 파일이 있어야 한다. 출력 위치는 판정에 쓰지 않는다
    (작업 폴더 밖이나 다운로드 폴더로 바로 변환해도 같은 게이트를 거친다).
    """
    real = src.resolve()  # 링크로 이름을 바꿔 붙인 초안은 실제 파일 기준으로 판단한다
    d = real.parent
    if real.name != "15_manuscript.md" or not (d / "_state.json").is_file():
        return "책 폴더(draft/{책})의 15_manuscript.md가 아닙니다. 초안 실측 같은 시험 변환은 --no-gate를 쓰세요."
    import book
    book.ROOT = d.parent.parent
    book.DRAFT, book.OUTPUT = book.ROOT / "draft", book.ROOT / "output"
    state = book.load_state(d)
    gate = book.publish_gate(d, state)
    return gate[1] if gate else None


def find_heading_pages(pdf, entries):
    """PDF 텍스트에서 차례 항목의 쪽을 찾는다. pdftotext가 없으면 None."""
    import shutil
    import subprocess
    if not shutil.which("pdftotext") or not pdf.is_file():
        return None
    res = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True, timeout=300)
    pages = [[C.despace(ln) for ln in pg.splitlines() if ln.strip()] for pg in res.stdout.split("\f")]
    titles = [C.despace(re.sub(r"\*\*|`|<[^>]+>", "", t)) for _, t in entries]

    def heads(pg):
        """쪽 머리의 줄들 (머리글 한 줄과 쪽 번호를 빼고 앞 세 줄)."""
        body = [ln for ln in pg if not ln.isdigit()]
        return body[1:4] if len(body) > 1 else body

    def starts_page(pg, title):
        # 제목이 쪽 머리에 '독립된 줄'로 있어야 한다. 길어서 줄바꿈된 제목은 첫 줄이 제목의 앞부분이다.
        # '부록의 프롬프트 템플릿…' 같은 본문 줄은 제목보다 길어 걸리지 않는다
        for ln in heads(pg):
            if ln == title or (len(ln) >= 4 and title.startswith(ln)) or (ln.startswith(title) and len(ln) <= len(title) + 3):
                return True
        return False

    toc_at = next((i for i, pg in enumerate(pages) if pg and any(ln == "차례" for ln in pg[:3])), None)
    if toc_at is None or not titles:
        return None
    toc_end = toc_at
    while toc_end + 1 < len(pages) and not starts_page(pages[toc_end + 1], titles[0]):
        toc_end += 1
    found, start = {}, toc_end + 1
    for idx, title in enumerate(titles):
        for pno in range(start, len(pages)):
            if title and starts_page(pages[pno], title):
                found[idx] = pno + 1
                start = pno
                break
    return found or None


def render_pdf(docx_path):
    """LibreOffice로 PDF를 만들고 쪽수를 돌려준다. soffice가 없거나 실패하면 None."""
    import shutil
    import subprocess
    import tempfile
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as prof:
        try:
            subprocess.run([exe, "--headless", f"-env:UserInstallation=file://{prof}", "--convert-to", "pdf",
                            "--outdir", str(docx_path.parent), str(docx_path)],
                           capture_output=True, timeout=600, check=False)
        except (OSError, subprocess.TimeoutExpired):
            return None
    pdf = docx_path.with_suffix(".pdf")
    if not pdf.is_file():
        return None
    data = pdf.read_bytes()
    counts = [int(x) for x in re.findall(rb"/Type\s*/Pages\b[^>]*?/Count\s+(\d+)", data)]
    return max(counts) if counts else len(re.findall(rb"/Type\s*/Page\b(?!s)", data))


if __name__ == "__main__":
    sys.exit(main())
